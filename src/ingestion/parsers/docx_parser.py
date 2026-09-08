"""DOCX parser.

DOCX has no page concept, so the heading path carries provenance. We walk the
document *body in order* rather than using ``doc.paragraphs`` then ``doc.tables``,
because losing the interleaving would detach a table from the heading that names
it — and in this archive the heading is usually the entity name.
"""

from __future__ import annotations

from pathlib import Path

from docx import Document as DocxDocument
from docx.table import Table as DocxTable
from docx.text.paragraph import Paragraph

from src.common.ids import table_id as make_table_id
from src.common.models import Block, Table
from src.ingestion.parsers.base import ParseResult


def _iter_body(document) -> list[Paragraph | DocxTable]:
    """Yield paragraphs and tables in true document order."""
    from docx.oxml.ns import qn

    body = document.element.body
    items: list[Paragraph | DocxTable] = []
    for child in body.iterchildren():
        if child.tag == qn("w:p"):
            items.append(Paragraph(child, document))
        elif child.tag == qn("w:tbl"):
            items.append(DocxTable(child, document))
    return items


def _heading_level(paragraph: Paragraph) -> int | None:
    style = (paragraph.style.name or "") if paragraph.style is not None else ""
    if style.startswith("Heading"):
        tail = style.replace("Heading", "").strip()
        return int(tail) if tail.isdigit() else 1
    if style == "Title":
        return 1
    return None


class DocxParser:
    def parse(self, path: Path, document_id: str, assets_dir: Path) -> ParseResult:
        result = ParseResult(page_count=None)
        document = DocxDocument(str(path))

        section: list[str] = []
        table_ordinal = 0

        for item in _iter_body(document):
            if isinstance(item, DocxTable):
                rows = [[cell.text.strip() for cell in row.cells] for row in item.rows]
                if not rows:
                    continue
                columns, body = rows[0], rows[1:]
                table = Table(
                    table_id=make_table_id(document_id, None, table_ordinal),
                    document_id=document_id,
                    page=None,
                    section_path=tuple(section),
                    caption=section[-1] if section else path.stem,
                    columns=columns,
                    rows=body,
                )
                result.tables.append(table)
                result.blocks.append(
                    Block(text=table.as_retrieval_text(), content_type="table",
                          section_path=tuple(section), table_id=table.table_id)
                )
                table_ordinal += 1
                continue

            text = item.text.strip()
            if not text:
                continue
            level = _heading_level(item)
            if level is not None:
                del section[level - 1 :]
                section.append(text)
                result.blocks.append(
                    Block(text=text, content_type="heading", section_path=tuple(section))
                )
            else:
                result.blocks.append(
                    Block(text=text, content_type="paragraph", section_path=tuple(section))
                )

        if not result.blocks:
            result.warnings.append("docx produced no text blocks")
        return result
